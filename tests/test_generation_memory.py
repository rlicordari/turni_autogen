import datetime as dt
import tempfile
import unittest
from pathlib import Path

import openpyxl

import generation_memory as gm


class GenerationMemoryTests(unittest.TestCase):
    def test_prior_usage_keeps_versions_but_counts_only_active_previous_dates(self):
        memory = gm.empty_memory()
        memory = gm.append_version(
            memory,
            version_id="v1",
            label="Settembre prima settimana",
            start_date=dt.date(2026, 9, 1),
            end_date=dt.date(2026, 9, 6),
            assignments={
                "2026-09-02": {"J": ["Zito"], "K": ["Allegra"]},
                "2026-09-04": {"J": ["Dattilo"]},
            },
            active=True,
        )
        memory = gm.append_version(
            memory,
            version_id="v2",
            label="Versione alternativa prima settimana",
            start_date=dt.date(2026, 9, 1),
            end_date=dt.date(2026, 9, 6),
            assignments={
                "2026-09-02": {"J": ["Zito"]},
            },
            active=False,
        )

        usage = gm.build_solver_prior_usage(
            memory,
            dt.date(2026, 9, 7),
            dt.date(2026, 9, 30),
        )

        self.assertEqual(len(memory["versions"]), 2)
        self.assertEqual(usage["counts"]["2026-09"]["Zito"]["J"], 1)
        self.assertEqual(usage["counts"]["2026-09"]["Dattilo"]["J"], 1)
        self.assertNotIn("v2", usage["versions_used"])
        self.assertEqual(usage["night_dates_by_doc"]["2026-09"]["Zito"], ["2026-09-02"])

    def test_prior_usage_overlapping_active_versions_use_latest_date_version(self):
        memory = gm.empty_memory()
        memory = gm.append_version(
            memory,
            version_id="old",
            label="Prima versione",
            start_date=dt.date(2026, 9, 1),
            end_date=dt.date(2026, 9, 6),
            assignments={"2026-09-02": {"J": ["Zito"]}},
            active=True,
        )
        memory = gm.append_version(
            memory,
            version_id="new",
            label="Seconda versione",
            start_date=dt.date(2026, 9, 1),
            end_date=dt.date(2026, 9, 6),
            assignments={"2026-09-02": {"J": ["Dattilo"]}},
            active=True,
        )

        usage = gm.build_solver_prior_usage(
            memory,
            dt.date(2026, 9, 7),
            dt.date(2026, 9, 30),
        )

        self.assertNotIn("Zito", usage["counts"]["2026-09"])
        self.assertEqual(usage["counts"]["2026-09"]["Dattilo"]["J"], 1)
        self.assertNotIn("old", usage["versions_used"])
        self.assertIn("new", usage["versions_used"])

    def test_prior_usage_uses_only_selected_versions_when_provided(self):
        memory = gm.empty_memory()
        memory = gm.append_version(
            memory,
            version_id="old",
            label="Prima versione",
            start_date=dt.date(2026, 9, 1),
            end_date=dt.date(2026, 9, 6),
            assignments={"2026-09-02": {"J": ["Zito"]}},
            active=True,
        )
        memory = gm.append_version(
            memory,
            version_id="new",
            label="Seconda versione",
            start_date=dt.date(2026, 9, 1),
            end_date=dt.date(2026, 9, 6),
            assignments={"2026-09-02": {"J": ["Dattilo"]}},
            active=True,
        )

        usage = gm.build_solver_prior_usage(
            memory,
            dt.date(2026, 9, 7),
            dt.date(2026, 9, 30),
            selected_version_ids={"old"},
        )

        self.assertEqual(usage["counts"]["2026-09"]["Zito"]["J"], 1)
        self.assertNotIn("Dattilo", usage["counts"]["2026-09"])
        self.assertEqual(usage["versions_used"], ["old"])

    def test_prior_usage_selected_inactive_version_is_used_explicitly(self):
        memory = gm.empty_memory()
        memory = gm.append_version(
            memory,
            version_id="inactive",
            label="Versione disattivata ma scelta",
            start_date=dt.date(2026, 9, 1),
            end_date=dt.date(2026, 9, 6),
            assignments={"2026-09-02": {"J": ["Zito"]}},
            active=False,
        )

        usage_without_selection = gm.build_solver_prior_usage(
            memory,
            dt.date(2026, 9, 7),
            dt.date(2026, 9, 30),
        )
        usage_with_selection = gm.build_solver_prior_usage(
            memory,
            dt.date(2026, 9, 7),
            dt.date(2026, 9, 30),
            selected_version_ids={"inactive"},
        )

        self.assertEqual(usage_without_selection["counts"], {})
        self.assertEqual(usage_with_selection["counts"]["2026-09"]["Zito"]["J"], 1)
        self.assertEqual(usage_with_selection["versions_used"], ["inactive"])

    def test_prior_usage_ignores_recupero_operationally(self):
        memory = gm.empty_memory()
        memory = gm.append_version(
            memory,
            version_id="v1",
            label="Recupero salvato",
            start_date=dt.date(2026, 9, 1),
            end_date=dt.date(2026, 9, 1),
            assignments={"2026-09-01": {"L": ["Recupero"], "T": ["Recupero"], "J": ["Recupero"]}},
            active=True,
        )

        usage = gm.build_solver_prior_usage(
            memory,
            dt.date(2026, 9, 7),
            dt.date(2026, 9, 30),
        )

        self.assertEqual(usage["counts"], {})
        self.assertEqual(usage["night_dates_by_doc"], {})

    def test_prior_usage_ignores_same_period_when_regenerating(self):
        memory = gm.empty_memory()
        memory = gm.append_version(
            memory,
            version_id="v1",
            label="Settembre prima settimana",
            start_date=dt.date(2026, 9, 1),
            end_date=dt.date(2026, 9, 6),
            assignments={"2026-09-02": {"J": ["Zito"]}},
            active=True,
        )

        usage = gm.build_solver_prior_usage(
            memory,
            dt.date(2026, 9, 1),
            dt.date(2026, 9, 6),
        )

        self.assertEqual(usage["counts"], {})
        self.assertEqual(usage["versions_used"], [])

    def test_prior_usage_includes_previous_month_nights_for_carryover(self):
        memory = gm.empty_memory()
        memory = gm.append_version(
            memory,
            version_id="v1",
            label="Fine settembre",
            start_date=dt.date(2026, 9, 30),
            end_date=dt.date(2026, 9, 30),
            assignments={"2026-09-30": {"J": ["Zito"]}},
            active=True,
        )

        usage = gm.build_solver_prior_usage(
            memory,
            dt.date(2026, 10, 1),
            dt.date(2026, 10, 31),
        )

        self.assertEqual(usage["counts"], {})
        self.assertEqual(usage["night_dates_by_doc"]["2026-10"]["Zito"], ["2026-09-30"])

    def test_prior_usage_ignores_generated_memory_for_finalized_months(self):
        memory = gm.empty_memory()
        memory = gm.append_version(
            memory,
            version_id="v1",
            label="Settembre provvisorio",
            start_date=dt.date(2026, 9, 30),
            end_date=dt.date(2026, 9, 30),
            assignments={"2026-09-30": {"J": ["Zito"]}},
            active=True,
        )

        usage = gm.build_solver_prior_usage(
            memory,
            dt.date(2026, 10, 1),
            dt.date(2026, 10, 31),
            finalized_months={"2026-09"},
        )

        self.assertEqual(usage["counts"], {})
        self.assertEqual(usage["night_dates_by_doc"], {})

    def test_parse_generated_xlsx_extracts_assignments(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "turni.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws["A1"] = "Data"
            ws["J1"] = "Notte"
            ws["K1"] = "Letto"
            ws["L1"] = "Padiglioni"
            ws["A2"] = dt.date(2026, 9, 2)
            ws["J2"] = "Zito"
            ws["K2"] = "Allegra"
            ws["L2"] = "Recupero"
            wb.save(path)

            assignments = gm.parse_generated_xlsx_assignments(path)

        self.assertEqual(assignments["2026-09-02"]["J"], ["Zito"])
        self.assertEqual(assignments["2026-09-02"]["K"], ["Allegra"])
        self.assertEqual(assignments["2026-09-02"]["L"], ["Recupero"])

    def test_memory_to_shift_history_counts_active_versions_like_finalized_history(self):
        memory = gm.empty_memory()
        memory = gm.append_version(
            memory,
            version_id="v1",
            label="Settembre prima settimana",
            start_date=dt.date(2026, 9, 1),
            end_date=dt.date(2026, 9, 6),
            assignments={
                "2026-09-02": {"J": ["Zito"], "K": ["Allegra"]},
                "2026-09-06": {"D": ["Crea"], "E": ["Crea"], "J": ["Dattilo"]},
            },
            active=True,
        )
        memory = gm.append_version(
            memory,
            version_id="v2",
            label="Alternativa inattiva",
            start_date=dt.date(2026, 9, 1),
            end_date=dt.date(2026, 9, 6),
            assignments={"2026-09-03": {"J": ["Zito"]}},
            active=False,
        )

        history = gm.memory_to_shift_history(
            memory,
            valid_doctors={"Allegra", "Crea", "Dattilo", "Zito"},
        )

        self.assertEqual(history["2026-09"]["Zito"]["J"]["total"], 1)
        self.assertEqual(history["2026-09"]["Dattilo"]["J"]["domeniche"], 1)
        self.assertEqual(history["2026-09"]["Crea"]["_festivi_DE_HI"], 1)
        self.assertEqual(history["2026-09"]["Crea"]["D"]["total"], 1)
        self.assertEqual(history["2026-09"]["Crea"]["E"]["total"], 1)
        self.assertNotIn("v2", history["2026-09"]["_meta"]["generation_versions_used"])

    def test_memory_to_shift_history_overlapping_active_versions_use_latest_date_version(self):
        memory = gm.empty_memory()
        memory = gm.append_version(
            memory,
            version_id="old",
            label="Prima versione",
            start_date=dt.date(2026, 9, 1),
            end_date=dt.date(2026, 9, 6),
            assignments={"2026-09-02": {"J": ["Zito"]}},
            active=True,
        )
        memory = gm.append_version(
            memory,
            version_id="new",
            label="Seconda versione",
            start_date=dt.date(2026, 9, 1),
            end_date=dt.date(2026, 9, 6),
            assignments={"2026-09-02": {"J": ["Dattilo"]}},
            active=True,
        )

        history = gm.memory_to_shift_history(memory)

        self.assertNotIn("Zito", history["2026-09"])
        self.assertEqual(history["2026-09"]["Dattilo"]["J"]["total"], 1)
        self.assertEqual(history["2026-09"]["_meta"]["generation_versions_used"], ["new"])

    def test_memory_to_shift_history_uses_only_selected_versions_when_provided(self):
        memory = gm.empty_memory()
        memory = gm.append_version(
            memory,
            version_id="old",
            label="Prima versione",
            start_date=dt.date(2026, 9, 1),
            end_date=dt.date(2026, 9, 6),
            assignments={"2026-09-02": {"J": ["Zito"]}},
            active=True,
        )
        memory = gm.append_version(
            memory,
            version_id="new",
            label="Seconda versione",
            start_date=dt.date(2026, 9, 1),
            end_date=dt.date(2026, 9, 6),
            assignments={"2026-09-02": {"J": ["Dattilo"]}},
            active=True,
        )

        history = gm.memory_to_shift_history(memory, selected_version_ids={"old"})

        self.assertEqual(history["2026-09"]["Zito"]["J"]["total"], 1)
        self.assertNotIn("Dattilo", history["2026-09"])
        self.assertEqual(history["2026-09"]["_meta"]["generation_versions_used"], ["old"])

    def test_memory_to_shift_history_selected_inactive_version_is_used_explicitly(self):
        memory = gm.empty_memory()
        memory = gm.append_version(
            memory,
            version_id="inactive",
            label="Versione disattivata ma scelta",
            start_date=dt.date(2026, 9, 1),
            end_date=dt.date(2026, 9, 6),
            assignments={"2026-09-02": {"J": ["Zito"]}},
            active=False,
        )

        history_without_selection = gm.memory_to_shift_history(memory)
        history_with_selection = gm.memory_to_shift_history(memory, selected_version_ids={"inactive"})

        self.assertEqual(history_without_selection, {})
        self.assertEqual(history_with_selection["2026-09"]["Zito"]["J"]["total"], 1)
        self.assertEqual(history_with_selection["2026-09"]["_meta"]["generation_versions_used"], ["inactive"])

    def test_memory_to_shift_history_ignores_recupero_operationally(self):
        memory = gm.empty_memory()
        memory = gm.append_version(
            memory,
            version_id="v1",
            label="Recupero salvato",
            start_date=dt.date(2026, 9, 1),
            end_date=dt.date(2026, 9, 1),
            assignments={"2026-09-01": {"L": ["Recupero"], "T": ["Recupero"], "J": ["Recupero"]}},
            active=True,
        )

        history = gm.memory_to_shift_history(memory)

        self.assertEqual(history, {})

    def test_memory_to_shift_history_respects_cutoff_and_finalized_months(self):
        memory = gm.empty_memory()
        memory = gm.append_version(
            memory,
            version_id="v1",
            label="Settembre",
            start_date=dt.date(2026, 9, 1),
            end_date=dt.date(2026, 9, 10),
            assignments={
                "2026-09-02": {"J": ["Zito"]},
                "2026-09-08": {"J": ["Dattilo"]},
            },
            active=True,
        )

        cutoff_history = gm.memory_to_shift_history(
            memory,
            start_date=dt.date(2026, 9, 7),
            end_date=dt.date(2026, 9, 30),
        )
        skipped_finalized = gm.memory_to_shift_history(
            memory,
            start_date=dt.date(2026, 10, 1),
            end_date=dt.date(2026, 10, 31),
            finalized_months={"2026-09"},
        )

        self.assertEqual(cutoff_history["2026-09"]["Zito"]["J"]["total"], 1)
        self.assertNotIn("Dattilo", cutoff_history["2026-09"])
        self.assertEqual(skipped_finalized, {})


if __name__ == "__main__":
    unittest.main()
